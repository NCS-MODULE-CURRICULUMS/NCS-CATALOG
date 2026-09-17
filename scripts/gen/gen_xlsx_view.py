# -*- coding: utf-8 -*-
"""
표준 강의 교안 xlsx → 정적 HTML 뷰어.

외부 라이브러리를 쓰지 않는 사이트 방침에 맞춰, 빌드 시점에 openpyxl 로 읽어
site.css 판면의 표로 미리 그려 둔다. 병합 셀은 colspan/rowspan 으로 옮기고,
원본의 배경색은 양식 범례(입력칸 / 라벨 / 자동계산 / 심사 중점)로 환산한다.

  python _shared/gen_xlsx_view.py
"""
import html
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

ORG = Path(__file__).resolve().parents[3]
SITE = ORG / "COURSE-MANAGEMENT"
SRC = SITE / "docs" / "파이썬_표준강의교안_샘플.xlsx"
OUT = SITE / "docs" / "표준강의교안-샘플.html"
DL = "파이썬_표준강의교안_샘플.xlsx"

# 원본 색 → 양식 범례. 작성안내 시트의 색상 범례를 그대로 따른다.
FILL = {
    "FFFFFF00": "in", "FFFFF2CC": "in", "FFFFE699": "in", "FFFFD966": "in",   # 노랑 = 입력칸
    "FFD9D9D9": "lb", "FFBFBFBF": "lb", "FFF2F2F2": "lb", "FFA6A6A6": "lb",   # 회색 = 라벨
    "FFC6E0B4": "gr", "FFE2EFDA": "gr", "FFA9D08E": "gr",                      # 연두 = 심사 중점
}


def cls_of(cell):
    f = cell.fill
    if f is None or f.fill_type != "solid":
        return ""
    rgb = getattr(f.fgColor, "rgb", None)
    if not isinstance(rgb, str):
        return ""
    return FILL.get(rgb.upper(), "")


def fmt(v):
    if v is None:
        return ""
    if isinstance(v, float):
        if v == int(v):
            return str(int(v))
        return f"{v:g}"
    return str(v).strip()


def sheet_html(ws):
    merged = {}          # (r,c) -> (rowspan, colspan)
    covered = set()
    for rng in ws.merged_cells.ranges:
        r1, c1, r2, c2 = rng.min_row, rng.min_col, rng.max_row, rng.max_col
        merged[(r1, c1)] = (r2 - r1 + 1, c2 - c1 + 1)
        for r in range(r1, r2 + 1):
            for c in range(c1, c2 + 1):
                if (r, c) != (r1, c1):
                    covered.add((r, c))

    # 실제로 값이 있는 마지막 행·열까지만
    maxr = maxc = 0
    for row in ws.iter_rows():
        for cell in row:
            if fmt(cell.value):
                maxr = max(maxr, cell.row)
                maxc = max(maxc, cell.column)
    if not maxr:
        return "<p class='memo'>내용이 없습니다.</p>"

    out = ['<div class="tscroll"><table class="xl">']
    for r in range(1, maxr + 1):
        cells = []
        for c in range(1, maxc + 1):
            if (r, c) in covered:
                continue
            cell = ws.cell(r, c)
            v = fmt(cell.value)
            rs, cs = merged.get((r, c), (1, 1))
            attr = ""
            if rs > 1:
                attr += f' rowspan="{rs}"'
            if cs > 1:
                attr += f' colspan="{cs}"'
            k = cls_of(cell)
            b = cell.font is not None and cell.font.bold
            if b and not k:
                k = "bd"
            if k:
                attr += f' class="{k}"'
            txt = html.escape(v).replace("\n", "<br>")
            cells.append(f"<td{attr}>{txt}</td>")
        if cells:
            out.append("<tr>" + "".join(cells) + "</tr>")
    out.append("</table></div>")
    return "\n".join(out)


def main():
    wb = openpyxl.load_workbook(SRC, data_only=True)
    names = wb.sheetnames
    tabs = "".join(
        f'<a href="#s{i}" data-i="{i}"{" class=\"act\"" if i == 0 else ""}>{html.escape(n)}</a>'
        for i, n in enumerate(names))
    panes = "".join(
        f'<section id="s{i}" class="pane"{"" if i == 0 else " hidden"}>'
        f'<h2 class="sub2">{html.escape(n)}</h2>{sheet_html(wb[n])}</section>'
        for i, n in enumerate(names))

    OUT.write_text(f"""<meta charset="utf-8"><title>표준 강의 교안 양식</title>
<link rel="stylesheet" href="../assets/site.css">
<script src="../assets/auth.js"></script>
<script>EXAM_AUTH.guard("../");</script>
<div class="top"><div class="tbar">
  <div class="brand"><a href="../index.html">과정관리</a><small>자료 / 표준 강의 교안 양식</small></div>
  <div class="tuser" id="tUser"><b id="tName"></b><button id="tOut" type="button">로그아웃</button></div>
</div>
<div class="nav" id="tabs">{tabs}</div>
</div>
<script>EXAM_AUTH.paintTop("../");</script>
<style>
  .sub2{{font-size:15px;margin:22px 0 8px;padding-bottom:6px;border-bottom:1px solid #000;font-weight:700}}
  .memo{{font-size:12.5px;line-height:1.7;color:#555}}
  .tscroll{{overflow-x:auto;-webkit-overflow-scrolling:touch}}
  table.xl{{font-size:12px;table-layout:auto}}
  table.xl td{{border:1px solid #999;padding:5px 7px;text-align:left;vertical-align:top;
               line-height:1.6;min-width:58px}}
  table.xl td.in{{background:#fffbe6}}        /* 작성자 입력칸 */
  table.xl td.lb{{background:#f0f0f0;font-weight:700;white-space:nowrap}}  /* 항목 라벨 */
  table.xl td.gr{{background:#eef6ee}}        /* 심사·인증 중점 */
  table.xl td.bd{{font-weight:700}}
  .lgd{{display:flex;gap:14px;flex-wrap:wrap;font-size:12px;margin:12px 0;
        border:1px solid #000;padding:9px 13px}}
  .lgd i{{display:inline-block;width:13px;height:13px;border:1px solid #999;
         vertical-align:-2px;margin-right:5px}}
  .bar{{display:flex;gap:6px;flex-wrap:wrap;margin:14px 0}}
  .bar a,.bar button{{border:1px solid #000;background:#fff;color:#000;padding:5px 14px;
              font-size:12.5px;cursor:pointer;font-family:inherit;text-decoration:none}}
  .bar a:hover,.bar button:hover{{background:#000;color:#fff}}
  [hidden]{{display:none}}
  @media print{{.top,.bar,.lgd{{display:none}} .pane[hidden]{{display:block}}}}
</style>
<div class="wrap">
<h1>표준 강의 교안 양식</h1>
<p class="sub">교과목 운영계획서 · 코리아AI아카데미 표준안 v1.0 · 과정평가형 · 일반국비(계좌제) · KDT 공통</p>

<div class="lgd">
  <span><i style="background:#fffbe6"></i>작성자 입력칸</span>
  <span><i style="background:#f0f0f0"></i>항목 라벨 — 수정하지 않음</span>
  <span><i style="background:#fff"></i>자동 계산 — 수정하지 않음</span>
  <span><i style="background:#eef6ee"></i>심사 · 인증 중점 확인</span>
</div>

<div class="bar">
  <a href="{DL}" download>xlsx 내려받기</a>
  <button id="bAll" type="button">전체 시트 펼치기</button>
  <button id="bPrint" type="button">인쇄 · PDF</button>
</div>

{panes}

<footer>원본: {DL} · 담당 정우균</footer>
</div>
<script>
(function(){{
  var tabs = document.getElementById('tabs'),
      links = [].slice.call(tabs.querySelectorAll('a')),
      panes = [].slice.call(document.querySelectorAll('.pane')),
      all = false;
  function show(i){{
    all = false;
    panes.forEach(function(p, j){{ p.hidden = (j !== i); }});
    links.forEach(function(a, j){{ a.className = (j === i) ? 'act' : ''; }});
    document.getElementById('bAll').textContent = '전체 시트 펼치기';
  }}
  links.forEach(function(a, i){{
    a.onclick = function(e){{ e.preventDefault(); show(i); window.scrollTo(0, 0); }};
  }});
  document.getElementById('bAll').onclick = function(){{
    all = !all;
    panes.forEach(function(p){{ p.hidden = !all && p.id !== 's0'; }});
    links.forEach(function(a, j){{ a.className = (!all && j === 0) ? 'act' : ''; }});
    this.textContent = all ? '한 시트만 보기' : '전체 시트 펼치기';
  }};
  document.getElementById('bPrint').onclick = function(){{ window.print(); }};
}})();
</script>
""", encoding="utf-8")
    print(f"{OUT.relative_to(SITE)}  시트 {len(names)}개 — {', '.join(names)}")
    print(f"크기 {OUT.stat().st_size // 1024}KB")


if __name__ == "__main__":
    main()
